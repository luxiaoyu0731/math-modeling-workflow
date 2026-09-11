"""Small explicit Markdown document adapters; unsupported blocks fail visibly."""
from __future__ import annotations
import re
from pathlib import Path


def blocks(text):
    lines=text.splitlines();i=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('```'):
            code=[]
            while i<len(lines) and not lines[i].strip().startswith('```'):code.append(lines[i]);i+=1
            if i==len(lines):raise ValueError('unclosed code fence')
            i+=1;yield 'code','\n'.join(code);continue
        if line.startswith('$$'):
            formula=line[2:]
            if formula.endswith('$$'):formula=formula[:-2]
            else:
                parts=[formula]
                while i<len(lines) and '$$' not in lines[i]:parts.append(lines[i]);i+=1
                if i==len(lines):raise ValueError('unclosed display formula')
                tail=lines[i].split('$$',1)
                if tail[1].strip():raise ValueError('display formula must occupy its own block')
                parts.append(tail[0]);i+=1;formula='\n'.join(parts)
            yield 'math',formula.strip();continue
        heading=re.match(r'^(#{1,6})\s+(.+)$',line)
        if heading:yield 'heading',(len(heading[1]),heading[2]);continue
        image=re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)',line)
        if image:yield 'image',(image[1],image[2]);continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
            parsed=[]
            for row in rows:
                cells=[c.strip() for c in row.strip('|').split('|')]
                if all(re.fullmatch(r':?-{3,}:?',c) for c in cells):continue
                parsed.append(cells)
            if not parsed or any(len(r)!=len(parsed[0]) for r in parsed):raise ValueError('inconsistent Markdown table columns')
            yield 'table',parsed;continue
        if '<!--' in line or line.startswith('<'):raise ValueError('HTML is unsupported in final prose')
        if line.startswith('\\[') or line.startswith('\\begin'):raise ValueError('use $$ blocks for display equations')
        paragraph=[line]
        while i<len(lines) and lines[i].strip() and not re.match(r'^(#|\$\$|```|\||!\[)',lines[i].strip()):paragraph.append(lines[i].strip());i+=1
        yield 'paragraph',' '.join(paragraph)


def make_docx(root,text,profile,path):
    from docx import Document
    from docx.shared import Pt,Cm,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from vendor.formula_omml import inline_formula_tokens,latex_to_omml,display_omml
    document=Document();sec=document.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(profile['margin_cm'])
    if sec.page_width<=sec.left_margin+sec.right_margin:raise ValueError('margins leave no usable page width')
    normal=document.styles['Normal'];normal.font.size=Pt(profile['font_pt'])
    normal.paragraph_format.line_spacing=profile['line_spacing']
    normal.paragraph_format.space_after=Pt(4)
    for name in ('Normal','Title','Heading 1','Heading 2','Heading 3','Caption'):
        style=document.styles[name];style.font.color.rgb=RGBColor(0,0,0)
        props=style.element.find(qn('w:pPr'))
        if props is not None:
            for border in list(props.findall(qn('w:pBdr'))):props.remove(border)
    if profile.get('font'):
        for name in ('Normal','Title','Heading 1','Heading 2','Heading 3','Caption'):
            style=document.styles[name];style.font.name=profile['font']
            fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
            for attr in list(fonts.attrib):
                if attr.endswith('Theme'):del fonts.attrib[attr]
            for attr in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+attr),profile['font'])
    counter={'inline_equations':0,'display_equations':0,'images':0,'tables':0}
    def plain(p,s):
        # Only basic emphasis is accepted; preserve punctuation and mathematical input.
        for idx,part in enumerate(re.split(r'(\*\*[^*]+\*\*)',s)):
            r=p.add_run(part[2:-2] if part.startswith('**') and part.endswith('**') else part)
            if part.startswith('**') and part.endswith('**'):r.bold=True
    def inline(p,s):
        offset=0
        for token in inline_formula_tokens(s):
            plain(p,s[offset:token.start]);p._p.append(latex_to_omml(token.latex));counter['inline_equations']+=1;offset=token.end
        plain(p,s[offset:])
    for kind,data in blocks(text):
        if kind=='heading':
            level,title=data;p=document.add_paragraph(style='Title' if level==1 else 'Heading '+str(min(level-1,3)));inline(p,title)
        elif kind=='paragraph':inline(document.add_paragraph(),data)
        elif kind=='code':
            p=document.add_paragraph();p.paragraph_format.line_spacing=1
            r=p.add_run(data);r.font.name='Courier New';r.font.size=Pt(9)
        elif kind=='math':document.add_paragraph()._p.append(display_omml(data));counter['display_equations']+=1
        elif kind=='image':
            from paper import local
            caption,relative=data;image=local(root,relative)
            p=document.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(image),width=sec.page_width-sec.left_margin-sec.right_margin)
            inline(document.add_paragraph(style='Caption'),caption);counter['images']+=1
        elif kind=='table':
            table=document.add_table(rows=0,cols=len(data[0]));table.style='Table Grid'
            for cells in data:
                row=table.add_row()
                for cell,value in zip(row.cells,cells):inline(cell.paragraphs[0],value)
            counter['tables']+=1
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    document.core_properties.author='';document.core_properties.last_modified_by=''
    document.save(path)
    return counter


def tex_escape(s):
    return ''.join({'\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_','{':r'\{','}':r'\}','~':r'\textasciitilde{}','^':r'\textasciicircum{}'}.get(c,c) for c in s)


def tex_inline(s):
    def plain(value):
        return ''.join(r'\textbf{'+tex_escape(part[2:-2])+'}' if part.startswith('**') and part.endswith('**') else tex_escape(part) for part in re.split(r'(\*\*[^*]+\*\*)',value))
    result=[];offset=0
    for m in re.finditer(r'(?<!\\)\$(?!\$)(.+?)(?<!\\)\$|\\\((.+?)\\\)',s):
        result.append(plain(s[offset:m.start()]));result.append('$'+(m[1] or m[2])+'$');offset=m.end()
    result.append(plain(s[offset:]));return ''.join(result)


def make_tex(root,text,profile,path,language):
    # Source paths are relative to project root; compile with that cwd.
    cls='ctexart' if language=='zh' else 'article'
    font_pt=profile['font_pt']
    margin=profile['margin_cm'];spacing=profile['line_spacing']
    options='a4paper,fontset=fandol' if language=='zh' else 'a4paper'
    parts=[r'\documentclass['+options+']{'+cls+'}',r'\usepackage[margin='+str(margin)+r'cm]{geometry}',
           r'\usepackage{amsmath,amssymb,graphicx,longtable,array,setspace}',r'\begin{document}',
           r'\fontsize{'+str(font_pt)+'}{'+str(font_pt*1.2)+r'}\selectfont',r'\setstretch{'+str(spacing)+'}']
    # A Word line-spacing multiplier is not asserted to have identical TeX metrics.
    for kind,data in blocks(text):
        if kind=='heading':
            level,title=data
            if level==1:parts.append(r'\begin{center}\Large\bfseries '+tex_inline(title)+r'\end{center}')
            else:parts.append(['',r'\section*{',r'\subsection*{',r'\subsubsection*{'][min(level-1,3)]+tex_inline(title)+'}')
        elif kind=='paragraph':parts.append(tex_inline(data)+'\n')
        elif kind=='math':parts.append('\\[\n'+data+'\n\\]')
        elif kind=='code':
            if '\\end{verbatim}' in data:raise ValueError('verbatim terminator in code')
            parts.append('\\begin{verbatim}\n'+data+'\n\\end{verbatim}')
        elif kind=='image':
            from paper import local
            caption,relative=data;local(root,relative)
            if any(c in relative for c in '{}%\\\n'):raise ValueError('unsupported character in TeX image path')
            parts.append(r'\begin{figure}[htbp]\centering\includegraphics[width=\linewidth]{\detokenize{'+relative+r'}}\caption{'+tex_inline(caption)+r'}\end{figure}')
        elif kind=='table':
            width=1/len(data[0]);columns='|'.join('p{'+str(round(width*.88,3))+r'\linewidth}' for _ in data[0])
            parts.append(r'\begin{longtable}{|'+columns+r'|}\hline')
            for row in data:parts.append(' & '.join(tex_inline(x) for x in row)+r'\\\hline')
            parts.append(r'\end{longtable}')
    parts.append(r'\end{document}');path.write_text('\n'.join(parts)+'\n',encoding='utf-8')
